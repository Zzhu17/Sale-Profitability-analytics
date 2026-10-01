import csv
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from pathlib import Path

from openpyxl import load_workbook

from b2b_domain.source_adapters import RawScalar


class Severity(StrEnum):
    CRITICAL = "critical"
    WARNING = "warning"
    INFORMATIONAL = "informational"


class Disposition(StrEnum):
    REJECT = "reject"
    QUARANTINE = "quarantine"
    FLAG = "flag"


class ValueState(StrEnum):
    MISSING = "missing"
    INVALID = "invalid"
    VALUE = "value"


@dataclass(frozen=True)
class TableSpec:
    name: str
    required_columns: frozenset[str]
    key_columns: tuple[str, ...]
    date_columns: tuple[str, ...] = ()
    decimal_columns: tuple[str, ...] = ()
    positive_columns: tuple[str, ...] = ()
    required_value_columns: tuple[str, ...] = ()
    warning_if_missing: tuple[tuple[str, str], ...] = ()


@dataclass(frozen=True)
class AuditIssue:
    severity: Severity
    code: str
    table: str
    row_number: int | None
    column: str | None
    detail: str

    @property
    def disposition(self) -> Disposition:
        if self.severity == Severity.CRITICAL:
            return Disposition.REJECT
        if self.severity == Severity.WARNING:
            return Disposition.QUARANTINE
        return Disposition.FLAG


@dataclass(frozen=True)
class AuditResult:
    table: str
    row_count: int
    issues: tuple[AuditIssue, ...]

    @property
    def critical_count(self) -> int:
        return sum(issue.severity == Severity.CRITICAL for issue in self.issues)


EXPECTED_TEMPLATE_SHEETS = frozenset(
    {
        "Start_Here",
        "Field_Guide",
        "Sales_Reps",
        "Customers",
        "Products",
        "Invoices",
        "Invoice_Items",
        "Returns",
        "Inventory",
        "Purchase_Batches",
    }
)


def _issue(
    code: str,
    spec: TableSpec,
    row_number: int | None,
    column: str | None,
    detail: str,
    severity: Severity = Severity.CRITICAL,
) -> AuditIssue:
    return AuditIssue(severity, code, spec.name, row_number, column, detail)


def audit_rows(
    rows: Iterable[Mapping[str, RawScalar]],
    spec: TableSpec,
    *,
    today: date | None = None,
) -> AuditResult:
    current_date = today or date.today()
    issues: list[AuditIssue] = []
    seen_keys: set[tuple[str, ...]] = set()
    row_count = 0

    for row_number, row in enumerate(rows, start=2):
        row_count += 1
        key = tuple(_text(row.get(column)) for column in spec.key_columns)
        for column, key_value in zip(spec.key_columns, key, strict=True):
            if not key_value:
                issues.append(
                    _issue("missing_key", spec, row_number, column, "Required key is blank")
                )
        if all(key):
            if key in seen_keys:
                issues.append(
                    _issue("duplicate_key", spec, row_number, ",".join(spec.key_columns), str(key))
                )
            seen_keys.add(key)

        for column in spec.required_value_columns:
            if _is_missing(row.get(column)):
                issues.append(
                    _issue("missing_value", spec, row_number, column, "Required value is blank")
                )
        for column, code in spec.warning_if_missing:
            if _is_missing(row.get(column)):
                issues.append(
                    _issue(
                        code,
                        spec,
                        row_number,
                        column,
                        "Value is unknown; it was not converted to zero",
                        Severity.WARNING,
                    )
                )

        for column in spec.date_columns:
            raw_value = _text(row.get(column))
            if _is_missing(row.get(column)):
                continue
            try:
                parsed = date.fromisoformat(raw_value)
            except ValueError:
                issues.append(
                    _issue("unparseable_date", spec, row_number, column, raw_value)
                )
                continue
            if parsed > current_date:
                issues.append(_issue("future_date", spec, row_number, column, raw_value))

        for column in spec.decimal_columns:
            raw_value = _text(row.get(column))
            if _is_missing(row.get(column)):
                continue
            try:
                decimal_value = Decimal(raw_value)
            except InvalidOperation:
                issues.append(
                    _issue("unparseable_decimal", spec, row_number, column, raw_value)
                )
                continue
            if column in spec.positive_columns and decimal_value <= 0:
                issues.append(
                    _issue("invalid_positive_value", spec, row_number, column, raw_value)
                )

    return AuditResult(spec.name, row_count, tuple(issues))


def _is_missing(value: RawScalar) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: RawScalar) -> str:
    return "" if value is None else str(value).strip()


def decimal_value(value: RawScalar) -> tuple[ValueState, Decimal | None]:
    if _is_missing(value):
        return ValueState.MISSING, None
    try:
        return ValueState.VALUE, Decimal(_text(value))
    except InvalidOperation:
        return ValueState.INVALID, None


def audit_csv(path: Path, spec: TableSpec, *, today: date | None = None) -> AuditResult:
    with path.open(encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        columns = frozenset(reader.fieldnames or ())
        missing = spec.required_columns - columns
        if missing:
            return AuditResult(
                spec.name,
                0,
                tuple(
                    _issue("missing_column", spec, 1, column, "Required column is absent")
                    for column in sorted(missing)
                ),
            )
        return audit_rows(reader, spec, today=today)


def audit_workbook_template(path: Path) -> tuple[AuditIssue, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    present = frozenset(workbook.sheetnames)
    spec = TableSpec(name="workbook_template", required_columns=frozenset(), key_columns=())
    return tuple(
        _issue("missing_sheet", spec, None, sheet, "Expected guide sheet is absent")
        for sheet in sorted(EXPECTED_TEMPLATE_SHEETS - present)
    )
