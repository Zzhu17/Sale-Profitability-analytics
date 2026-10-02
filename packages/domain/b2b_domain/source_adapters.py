import csv
import hashlib
from collections.abc import Iterator, Mapping
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Protocol

from openpyxl import load_workbook

type RawScalar = str | int | float | bool | None


class SourceAdapterError(ValueError):
    pass


@dataclass(frozen=True)
class RawRecord:
    source_name: str
    original_row_number: int
    values: Mapping[str, RawScalar]


@dataclass(frozen=True)
class SourceSchema:
    source_name: str
    headers: tuple[str, ...]


class SourceAdapter(Protocol):
    def schemas(self) -> Iterator[SourceSchema]: ...

    def records(self) -> Iterator[RawRecord]: ...


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _headers(values: tuple[object, ...], source_name: str) -> tuple[str, ...]:
    headers = tuple("" if value is None else str(value).strip() for value in values)
    if not headers or any(not header for header in headers):
        raise SourceAdapterError(f"{source_name}: header contains a blank column name")
    if len(headers) != len(set(headers)):
        raise SourceAdapterError(f"{source_name}: header contains duplicate column names")
    return headers


def _json_scalar(value: object) -> RawScalar:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


class CsvSourceAdapter:
    def __init__(self, path: Path) -> None:
        self.path = path

    def schemas(self) -> Iterator[SourceSchema]:
        with self.path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source)
            try:
                headers = _headers(tuple(next(reader)), self.path.stem)
            except StopIteration as error:
                raise SourceAdapterError(f"{self.path.name}: source is empty") from error
            yield SourceSchema(self.path.stem, headers)

    def records(self) -> Iterator[RawRecord]:
        with self.path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source)
            try:
                headers = _headers(tuple(next(reader)), self.path.stem)
            except StopIteration:
                return
            for row_number, row in enumerate(reader, start=2):
                if not any(value != "" for value in row):
                    continue
                if len(row) != len(headers):
                    raise SourceAdapterError(
                        f"{self.path.name}:{row_number}: expected {len(headers)} fields, "
                        f"found {len(row)}"
                    )
                yield RawRecord(self.path.stem, row_number, dict(zip(headers, row, strict=True)))


class ExcelSourceAdapter:
    def __init__(
        self,
        path: Path,
        header_rows: Mapping[str, int] | None = None,
        included_sheets: frozenset[str] | None = None,
    ) -> None:
        self.path = path
        self.header_rows = header_rows or {}
        self.included_sheets = included_sheets

    def schemas(self) -> Iterator[SourceSchema]:
        workbook = load_workbook(self.path, read_only=True, data_only=False)
        try:
            for worksheet in workbook.worksheets:
                if (
                    self.included_sheets is not None
                    and worksheet.title not in self.included_sheets
                ):
                    continue
                header_row = self.header_rows.get(worksheet.title, 1)
                rows = worksheet.iter_rows(values_only=True)
                for _ in range(header_row - 1):
                    next(rows, None)
                header_values = next(rows, None)
                if header_values is None or not any(value is not None for value in header_values):
                    continue
                trimmed_headers = list(header_values)
                while trimmed_headers and trimmed_headers[-1] is None:
                    trimmed_headers.pop()
                headers = _headers(tuple(trimmed_headers), worksheet.title)
                yield SourceSchema(worksheet.title, headers)
        finally:
            workbook.close()

    def records(self) -> Iterator[RawRecord]:
        workbook = load_workbook(self.path, read_only=True, data_only=False)
        try:
            for worksheet in workbook.worksheets:
                if (
                    self.included_sheets is not None
                    and worksheet.title not in self.included_sheets
                ):
                    continue
                header_row = self.header_rows.get(worksheet.title, 1)
                rows = worksheet.iter_rows(values_only=True)
                for _ in range(header_row - 1):
                    next(rows, None)
                header_values = next(rows, None)
                if header_values is None or not any(value is not None for value in header_values):
                    continue
                trimmed_headers = list(header_values)
                while trimmed_headers and trimmed_headers[-1] is None:
                    trimmed_headers.pop()
                headers = _headers(tuple(trimmed_headers), worksheet.title)
                for row_number, row in enumerate(rows, start=header_row + 1):
                    if not any(value is not None for value in row):
                        continue
                    values = tuple(_json_scalar(value) for value in row[: len(headers)])
                    if any(value is not None for value in row[len(headers) :]):
                        raise SourceAdapterError(
                            f"{worksheet.title}:{row_number}: data exists beyond the header"
                        )
                    yield RawRecord(
                        worksheet.title,
                        row_number,
                        dict(zip(headers, values, strict=True)),
                    )
        finally:
            workbook.close()


def adapter_for(path: Path) -> SourceAdapter:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return CsvSourceAdapter(path)
    if suffix in {".xlsx", ".xlsm"}:
        return ExcelSourceAdapter(path)
    raise SourceAdapterError(f"Unsupported source type: {suffix or '<none>'}")
